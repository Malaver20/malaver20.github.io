import os
import sys

def create_docx():
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    doc = docx.Document()

    # Margenes
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    COLOR_PRIMARY = RGBColor(216, 27, 96)     # Magenta #D81B60
    COLOR_TEAL = RGBColor(13, 148, 136)       # Teal #0D9488
    COLOR_NAVY = RGBColor(10, 17, 40)         # Navy #0A1128
    COLOR_TEXT = RGBColor(30, 41, 59)
    COLOR_MUTED = RGBColor(100, 116, 139)
    COLOR_AMBER = RGBColor(217, 119, 6)

    # Header
    logo_path = os.path.abspath('img/logo_facturanex_v3.png')
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(2.2)
    header_table.columns[1].width = Inches(4.8)

    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)

    if os.path.exists(logo_path):
        p_logo = cell_left.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_logo.add_run().add_picture(logo_path, width=Inches(1.8))
    else:
        p_logo = cell_left.paragraphs[0]
        r_brand = p_logo.add_run("FacturaNex SAS")
        r_brand.font.bold = True
        r_brand.font.size = Pt(18)
        r_brand.font.color.rgb = COLOR_PRIMARY

    p_right = cell_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sub = p_right.add_run("FACTURANEX SAS — NIT 901.555.123-4\n")
    r_sub.font.bold = True
    r_sub.font.size = Pt(8.5)
    r_sub.font.color.rgb = COLOR_NAVY
    
    r_cat = p_right.add_run("💊 PLANES COPI PARA COPIDROGUISTAS")
    r_cat.font.bold = True
    r_cat.font.size = Pt(9)
    r_cat.font.color.rgb = COLOR_TEAL

    doc.add_paragraph()

    # Titulo Principal
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BROCHURE COMERCIAL — PLANES COPI DROGUERÍAS\n")
    r_title.font.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = COLOR_PRIMARY

    r_integ = p_title.add_run("🔗 Integración Garantizada con Dominium Plus y Conexión Pfarma")
    r_integ.font.bold = True
    r_integ.font.size = Pt(10)
    r_integ.font.color.rgb = COLOR_TEAL

    p_desc = doc.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_desc = p_desc.add_run("Facturación Electrónica Ilimitada, POS Electrónico con Tirillas de Mostrador, Nómina y Recepción de Documentos.")
    r_desc.font.size = Pt(10)
    r_desc.font.color.rgb = COLOR_MUTED

    doc.add_paragraph()

    # Planes COPI
    planes = [
        {
            "nombre": "1. CopiFactura",
            "sub": "Facturación Electrónica DIAN Ilimitada + POS Mostrador",
            "precio": "$350.000 / año",
            "opcion": "📌 Resolución Adicional: +$160.000 / año",
            "items": [
                "• Facturación Electrónica DIAN Ilimitada",
                "• POS Electrónico (Tirilla Térmica Mostrador < 2 seg)",
                "• Integración con lectores de código de barras e impresoras térmicas",
                "• Habilitación ante la DIAN y Firma Digital Incluida"
            ]
        },
        {
            "nombre": "2. CopiNomina",
            "sub": "Nómina Electrónica Legal hasta 8 Empleados",
            "precio": "$350.000 / año",
            "opcion": "Diseñado para Regentes, Auxiliares y Domiciliarios",
            "items": [
                "• Transmisión ilimitada de Nómina mensual a la DIAN",
                "• Hasta 8 Empleados incluidos en el sistema",
                "• Desprendibles de pago en PDF enviados por WhatsApp / Email",
                "• Control de devengados, deducciones y prestaciones tributarias"
            ]
        },
        {
            "nombre": "3. CopiRecepcion",
            "sub": "Recepción Ilimitada de Facturas de Proveedores",
            "precio": "$260.000 / año",
            "opcion": "Buzón RADIAN y Deducción 100% DIAN",
            "items": [
                "• Acuse de recibo automático a laboratorios y depósitos (Copidrogas, etc.)",
                "• Eventos DIAN: Recepción de bienes/servicios y Aceptación",
                "• Garantiza la deducción del 100% de tus compras en la declaración de renta",
                "• Almacenamiento seguro de archivos XML"
            ]
        },
        {
            "nombre": "4. CopiDocumentoSoporte",
            "sub": "Documentos Soporte Ilimitados",
            "precio": "$260.000 / año",
            "opcion": "Legalización de compras a no obligados a facturar",
            "items": [
                "• Emisión de Documento Soporte Electrónico Ilimitado",
                "• Transmisión directa a la DIAN",
                "• Legalización de compras menores, fletes y servicios",
                "• Conexión automática con la plataforma cloud"
            ]
        }
    ]

    for p in planes:
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
        cell._tc.get_or_add_tcPr().append(shd)

        p_cell = cell.paragraphs[0]
        r_pn = p_cell.add_run(f"{p['nombre']} — ")
        r_pn.font.bold = True
        r_pn.font.size = Pt(11.5)
        r_pn.font.color.rgb = COLOR_PRIMARY

        r_pr = p_cell.add_run(f"{p['precio']}\n")
        r_pr.font.bold = True
        r_pr.font.size = Pt(12)
        r_pr.font.color.rgb = COLOR_NAVY

        r_ps = p_cell.add_run(f"{p['sub']} | {p['opcion']}\n")
        r_ps.font.italic = True
        r_ps.font.size = Pt(9)
        r_ps.font.color.rgb = COLOR_MUTED

        for item in p['items']:
            p_i = cell.add_paragraph()
            p_i.paragraph_format.space_before = Pt(1)
            p_i.paragraph_format.space_after = Pt(1)
            r_i = p_i.add_run(item)
            r_i.font.size = Pt(9)

        doc.add_paragraph()

    # COMBO 3 PLANES POR $860.000
    combo_table = doc.add_table(rows=1, cols=1)
    combo_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_combo = combo_table.cell(0, 0)
    shd_combo = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0A1128"/>')
    cell_combo._tc.get_or_add_tcPr().append(shd_combo)

    p_c = cell_combo.paragraphs[0]
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r_c1 = p_c.add_run("🎁 COMBO 3 PLANES COPI — ¡PREGUNTA POR TU OBSEQUIO!\n")
    r_c1.font.bold = True
    r_c1.font.size = Pt(13)
    r_c1.font.color.rgb = RGBColor(217, 119, 6) # Amber/Gold

    r_c2 = p_c.add_run("Elige 3 módulos para tu Droguería con descuento especial:\n")
    r_c2.font.size = Pt(9.5)
    r_c2.font.color.rgb = RGBColor(203, 213, 225)

    r_c3 = p_c.add_run("Inversión Especial Combo: 3 PLANES POR $860.000 / año\n")
    r_c3.font.bold = True
    r_c3.font.size = Pt(14)
    r_c3.font.color.rgb = RGBColor(255, 64, 129)

    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(10)
    r_f = p_foot.add_run("📲 WhatsApp / Teléfono: 315 515 55 12   |   🌐 facturanex.com")
    r_f.font.bold = True
    r_f.font.size = Pt(10)
    r_f.font.color.rgb = COLOR_NAVY

    doc.save("Brochure_Planes_COPI_Droguerias.docx")
    print("WORD updated successfully.")

def create_pdf():
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, HRFlowable
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    except ImportError:
        print("Reportlab missing")
        return

    doc = SimpleDocTemplate("Brochure_Planes_COPI_Droguerias.pdf", pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    story = []
    styles = getSampleStyleSheet()

    COLOR_PRIMARY = colors.HexColor('#D81B60')
    COLOR_TEAL = colors.HexColor('#0D9488')
    COLOR_NAVY = colors.HexColor('#0A1128')
    COLOR_TEXT = colors.HexColor('#1E293B')
    COLOR_MUTED = colors.HexColor('#64748B')
    COLOR_BG = colors.HexColor('#F8FAFC')

    title_style = ParagraphStyle('T1', fontName='Helvetica-Bold', fontSize=16, leading=20, textColor=COLOR_PRIMARY, alignment=TA_CENTER)
    sub_style = ParagraphStyle('T2', fontName='Helvetica', fontSize=9, leading=12, textColor=COLOR_MUTED, alignment=TA_CENTER)

    logo_file = os.path.abspath('img/logo_facturanex_v3.png')
    img_logo = Image(logo_file, width=120, height=38) if os.path.exists(logo_file) else Paragraph("<b>FacturaNex SAS</b>", styles['Normal'])

    head_right = Paragraph("<font color='#0A1128' size=8><b>FACTURANEX SAS — NIT 901.555.123-4</b></font><br/><font color='#0D9488' size=8.5><b>💊 BROCHURE PLANES COPI — COPIDROGUISTAS</b></font>", ParagraphStyle('R1', alignment=TA_RIGHT))
    header_table = Table([[img_logo, head_right]], colWidths=[200, 340])
    header_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'MIDDLE')]))

    story.append(header_table)
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1, color=COLOR_TEAL, spaceBefore=0, spaceAfter=8))

    story.append(Paragraph("BROCHURE COMERCIAL — PLANES COPI DROGUERÍAS", title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<font color='#0D9488'><b>🔗 Integración Garantizada con Dominium Plus y Conexión Pfarma</b></font>", ParagraphStyle('Integ', alignment=TA_CENTER, fontSize=9.5)))
    story.append(Spacer(1, 4))
    story.append(Paragraph("Facturación Electrónica Ilimitada, POS Electrónico con Tirilla Térmica de Mostrador, Nómina y Recepción de Documentos.", sub_style))
    story.append(Spacer(1, 10))

    planes = [
        {
            "name": "1. CopiFactura ($350.000 / año)",
            "sub": "Facturación Electrónica DIAN Ilimitada + POS Mostrador",
            "opc": "Resolución Adicional: +$160.000 / año",
            "bullets": "• Facturación Electrónica DIAN Ilimitada<br/>• POS Electrónico (Tirilla Térmica Mostrador &lt; 2 seg)<br/>• Compatible con impresoras térmicas (58mm/80mm) y lector de barras<br/>• Habilitación ante la DIAN y Firma Digital Incluida"
        },
        {
            "name": "2. CopiNomina ($350.000 / año)",
            "sub": "Nómina Electrónica Legal hasta 8 Empleados",
            "opc": "Regentes, Auxiliares y Domiciliarios",
            "bullets": "• Transmisión ilimitada de Nómina mensual a la DIAN<br/>• Hasta 8 Empleados incluidos<br/>• Desprendibles de pago en PDF por WhatsApp / Email<br/>• Control de devengados, deducciones y prestaciones"
        },
        {
            "name": "3. CopiRecepcion ($260.000 / año)",
            "sub": "Recepción de Documentos Ilimitada (Proveedores / Laboratorios)",
            "opc": "Acuse RADIAN y Deducción 100% DIAN",
            "bullets": "• Acuse de recibo automático a laboratorios y depósitos (Copidrogas, etc.)<br/>• Eventos DIAN: Recepción de bienes/servicios y Aceptación<br/>• Deducción 100% legal de compras en la declaración de renta<br/>• Almacenamiento seguro de archivos XML"
        },
        {
            "name": "4. CopiDocumentoSoporte ($260.000 / año)",
            "sub": "Documentos Soporte Ilimitados",
            "opc": "Compras a no obligados a facturar",
            "bullets": "• Emisión de Documento Soporte Electrónico Ilimitado<br/>• Transmisión directa a la DIAN<br/>• Legalización de compras menores, fletes y servicios<br/>• Conexión automática con la plataforma cloud"
        }
    ]

    for p in planes:
        cell_c = [
            Paragraph(f"<b><font color='#D81B60' size=10.5>{p['name']}</font></b> &nbsp;&nbsp; <font color='#D97706' size=8><b>[{p['opc']}]</b></font>", styles['Normal']),
            Paragraph(f"<i><font color='#64748B' size=8>{p['sub']}</font></i>", styles['Normal']),
            Spacer(1, 2),
            Paragraph(p['bullets'], ParagraphStyle('B1', fontName='Helvetica', fontSize=8, leading=10, textColor=COLOR_TEXT))
        ]
        t = Table([[cell_c]], colWidths=[540])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), COLOR_BG),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(t)
        story.append(Spacer(1, 6))

    combo_c = [
        Paragraph("<font color='#D97706' size=11><b>🎁 COMBO 3 PLANES COPI — ¡PREGUNTA POR TU OBSEQUIO!</b></font>", ParagraphStyle('C1', alignment=TA_CENTER)),
        Spacer(1, 2),
        Paragraph("<font color='#FFFFFF' size=12><b>Inversión Especial: 3 PLANES POR $860.000 / año</b></font>", ParagraphStyle('C2', alignment=TA_CENTER)),
        Paragraph("<font color='#CBD5E1' size=8.5>Elige 3 módulos para tu Droguería y reclama tu obsequio comercial</font>", ParagraphStyle('C3', alignment=TA_CENTER))
    ]
    t_combo = Table([[combo_c]], colWidths=[540])
    t_combo.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), COLOR_NAVY),
        ('BOX', (0,0), (-1,-1), 1.5, COLOR_PRIMARY),
        ('PADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t_combo)
    story.append(Spacer(1, 8))

    story.append(Paragraph("📲 <b>WhatsApp / Teléfono: 315 515 55 12</b> &nbsp;&nbsp;|&nbsp;&nbsp; 🌐 <b>facturanex.com</b>", ParagraphStyle('ContactSub', alignment=TA_CENTER, fontSize=9.5)))

    doc.build(story)
    print("PDF updated successfully.")

if __name__ == "__main__":
    create_docx()
    create_pdf()
